from html.parser import HTMLParser
from io import BytesIO
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import load_workbook

from app import app
from generator import generate_campaign
from google_ads_editor_export import CSV_PATH, README_PATH
from test_generator import submission
from validation import SERVICES


class FormParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.forms = []
        self.current = None

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == 'form':
            self.current = {'attrs': attrs, 'inputs': [], 'buttons': []}
            self.forms.append(self.current)
        elif tag == 'input' and self.current is not None:
            self.current['inputs'].append(attrs)
        elif tag == 'button' and self.current is not None:
            self.current['buttons'].append(attrs)

    def handle_endtag(self, tag):
        if tag == 'form':
            self.current = None


class GoogleAdsEditorRouteTests(unittest.TestCase):
    def setUp(self):
        self.client = app.test_client()

    def test_valid_download_headers_and_archive(self):
        form = submission(
            ['Emergency Plumbing', 'Boiler Repair'], website_url='https://example.com',
            business_name='Fictional Home Services',
        )
        with patch('app.generate_campaign', wraps=generate_campaign) as generate:
            response = self.client.post('/export/google-ads-editor', data=form)
            self.assertEqual(generate.call_count, 1)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, 'application/zip')
        self.assertIn('fictional-home-services-google-ads-editor.zip',
                      response.headers['Content-Disposition'])
        self.assertNotIn('Set-Cookie', response.headers)
        with ZipFile(BytesIO(response.data)) as archive:
            self.assertEqual(archive.namelist(), [CSV_PATH, README_PATH])
        self.assertEqual(self.client.get('/export/google-ads-editor').status_code, 405)

    def test_missing_and_invalid_website_retains_input(self):
        cases = ['', 'http://example.com', 'https://example.com/path', 'javascript:alert(1)']
        for value in cases:
            with self.subTest(value=value):
                form = submission(website_url=value)
                with patch('app.generate_campaign') as generate:
                    response = self.client.post('/export/google-ads-editor', data=form)
                    generate.assert_not_called()
                self.assertEqual(response.status_code, 400)
                self.assertEqual(response.mimetype, 'text/html')
                self.assertNotIn('Content-Disposition', response.headers)
                self.assertFalse(response.data.startswith(b'PK'))
                self.assertIn(f'value="{value}"', response.get_data(as_text=True))

    def test_formula_error_names_field_without_echoing_raw_value_in_error(self):
        form = submission(business_name='=SUM(1,2)', website_url='https://example.com')
        response = self.client.post('/export/google-ads-editor', data=form)
        self.assertEqual(response.status_code, 400)
        html = response.get_data(as_text=True)
        self.assertIn('Unsafe spreadsheet content was found in Headline', html)
        self.assertNotIn('=SUM(1,2)', html)

    def test_results_control_and_stateless_transport(self):
        for website, disabled in [('', True), ('https://example.com', False)]:
            with self.subTest(website=website):
                response = self.client.post('/generate', data=submission(SERVICES, website_url=website))
                self.assertEqual(response.status_code, 200)
                html = response.get_data(as_text=True)
                parser = FormParser()
                parser.feed(html)
                editor_forms = [form for form in parser.forms
                                if form['attrs'].get('action') == '/export/google-ads-editor']
                self.assertEqual(len(editor_forms), 1)
                form = editor_forms[0]
                values = {item.get('name'): item.get('value') for item in form['inputs']
                          if item.get('name') != 'services'}
                self.assertEqual(values['website_url'], website)
                self.assertEqual('disabled' in form['buttons'][0], disabled)
                for rendered in parser.forms:
                    names = [item.get('name') for item in rendered['inputs']]
                    if rendered['attrs'].get('action') in [
                            '/export', '/export/landing-pages', '/export/google-ads-editor'] or \
                            rendered['attrs'].get('action', '').startswith('/preview/'):
                        self.assertIn('website_url', names)
                self.assertIn('Not provided' if not website else website, html)

    def test_existing_excel_and_landing_page_exports_remain_intact(self):
        form = submission(
            ['Drain Cleaning'], website_url='https://example.com',
            business_phone='6175550123', primary_goal='Phone calls and quote/service requests',
        )
        excel = self.client.post('/export', data=form)
        self.assertEqual(excel.status_code, 200)
        book = load_workbook(BytesIO(excel.data))
        self.assertEqual(book.sheetnames, [
            'Overview', 'Campaign Structure', 'Keywords', 'Negative Keywords',
            'Ads', 'Budget', 'Landing Pages', 'Conversion Tracking',
        ])
        landing = self.client.post('/export/landing-pages', data=form)
        self.assertEqual(landing.status_code, 200)
        with ZipFile(BytesIO(landing.data)) as archive:
            self.assertEqual(archive.namelist(), [
                'index.html', 'drain-cleaning.html', 'assets/landing_page.css', 'README.txt',
            ])
            page = archive.read('drain-cleaning.html').decode('utf-8')
        self.assertIn('tel:+16175550123', page)
        self.assertIn('does not send or save requests', page)


if __name__ == '__main__':
    unittest.main()
