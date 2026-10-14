from html.parser import HTMLParser
from io import BytesIO
from itertools import combinations, product
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from app import app
from generator import generate_campaign
from landing_pages import generate_landing_pages
from rules_library import RULES
from validation import SERVICES, validate_intake
from test_generator import submission


class PageParser(HTMLParser):
    """Inspect controls/links structurally, including escaped hidden values."""
    def __init__(self):
        super().__init__()
        self.forms = []
        self.current_form = None
        self.nested_forms = False
        self.tags = []
        self.links = []
        self.ids = set()

    def handle_starttag(self, tag, attributes):
        attrs = dict(attributes)
        self.tags.append((tag, attrs))
        if 'id' in attrs:
            self.ids.add(attrs['id'])
        if tag == 'a':
            self.links.append(attrs)
        if tag == 'form':
            if self.current_form is not None:
                self.nested_forms = True
            self.current_form = {'attributes': attrs, 'inputs': []}
            self.forms.append(self.current_form)
        if tag == 'input' and self.current_form is not None:
            self.current_form['inputs'].append(attrs)

    def handle_endtag(self, tag):
        if tag == 'form':
            self.current_form = None


def parse(html):
    parser = PageParser()
    parser.feed(html)
    return parser


class LandingPageRouteTests(unittest.TestCase):
    def setUp(self):
        self.client = app.test_client()

    def path(self, service):
        return '/preview' + RULES[service]['path']

    def test_all_services_goals_and_availability(self):
        for service, goal, availability in product(SERVICES,
                ['Phone calls', 'Quote/service requests', 'Phone calls and quote/service requests'], ['Yes', 'No']):
            with self.subTest(service=service, goal=goal, availability=availability):
                form = submission([service], primary_goal=goal, emergency_availability=availability, state='Massachusetts')
                response = self.client.post(self.path(service), data=form)
                self.assertEqual(response.status_code, 200)
                html = response.get_data(as_text=True)
                parser = parse(html)
                self.assertIn('Local preview — not published.', html)
                self.assertIn(service, html)
                self.assertIn('Boston, Massachusetts', html)
                self.assertIn('Placeholder: add verified license information, if applicable.', html)
                self.assertIn('Placeholder: add the verified business phone number.', html)
                calls = goal != 'Quote/service requests'
                forms = goal != 'Phone calls'
                self.assertEqual('phone-placeholder' in parser.ids, calls)
                self.assertEqual('request-service' in parser.ids, forms)
                hrefs = [link.get('href') for link in parser.links]
                self.assertEqual('#phone-placeholder' in hrefs, calls)
                self.assertEqual('#request-service' in hrefs, forms)
                self.assertFalse(any(href and href.startswith('tel:') for href in hrefs))
                if availability == 'No':
                    for claim in ['24/7', '24 hour', 'day or night', 'overnight', 'round-the-clock']:
                        self.assertNotIn(claim, html.lower())
                else:
                    self.assertIn('24/7 emergency service', html)
                self.assertFalse(any(tag == 'script' for tag, attrs in parser.tags))
                self.assertIn('/static/landing_page.css', html)
                self.assertNotIn('/static/style.css', html)
                # Every internal CTA is a real anchor with a matching section.
                for href in hrefs:
                    if href and href.startswith('#'):
                        self.assertIn(href[1:], parser.ids)

    def test_disabled_mockup_cannot_submit(self):
        for goal in ['Quote/service requests', 'Phone calls and quote/service requests']:
            with self.subTest(goal=goal):
                html = self.client.post(self.path('Drain Cleaning'),
                    data=submission(['Drain Cleaning'], primary_goal=goal)).get_data(as_text=True)
                parser = parse(html)
                self.assertEqual(parser.forms, [])
                controls = [(tag, attrs) for tag, attrs in parser.tags if tag in ['input', 'textarea', 'button']]
                self.assertEqual(len(controls), 4)
                for tag, attrs in controls:
                    self.assertIn('disabled', attrs)
                    self.assertNotIn('name', attrs)
                    if tag == 'button':
                        self.assertEqual(attrs['type'], 'button')
                self.assertIn('This form is a preview and does not send or save requests.', html)
        calls = self.client.post(self.path('Drain Cleaning'), data=submission(['Drain Cleaning'])).get_data(as_text=True)
        self.assertFalse(any(tag in ['input', 'textarea', 'button', 'form'] for tag, attrs in parse(calls).tags))

    def test_known_gets_redirect_and_unknown_slugs_reject(self):
        for service in SERVICES:
            response = self.client.get(self.path(service))
            self.assertEqual(response.status_code, 302)
            self.assertEqual(response.headers['Location'], '/')
        for slug in ['unknown', 'landing_page.html', 'results.html']:
            self.assertEqual(self.client.get('/preview/' + slug).status_code, 404)
            self.assertEqual(self.client.post('/preview/' + slug, data=submission()).status_code, 404)
        self.assertEqual(self.client.post(self.path('Boiler Repair'), data=submission(['Emergency Plumbing'])).status_code, 404)

    def test_invalid_input_retains_values_and_precedes_slug_check(self):
        for url in [self.path('Emergency Plumbing'), '/preview/unknown']:
            for form in [submission([], monthly_budget='2500.01'), submission(monthly_budget='-10'),
                         submission(primary_goal='bad'), submission(emergency_availability='Maybe')]:
                with self.subTest(url=url, form=form):
                    response = self.client.post(url, data=form)
                    self.assertEqual(response.status_code, 400)
                    html = response.get_data(as_text=True)
                    self.assertIn('Google Ads Campaign Generator', html)
                    parser = parse(html)
                    inputs = {attrs['name']: attrs.get('value') for attrs in parser.forms[0]['inputs'] if 'name' in attrs}
                    self.assertEqual(inputs['business_name'], form['business_name'])
                    self.assertEqual(inputs['city'], form['city'])
                    self.assertEqual(inputs['monthly_budget'], form['monthly_budget'])

    def test_only_requested_selected_page_is_generated(self):
        with patch('app.generate_landing_pages', wraps=generate_landing_pages) as generate:
            response = self.client.post(self.path('Boiler Repair'), data=submission(SERVICES))
            self.assertEqual(response.status_code, 200)
            self.assertEqual(generate.call_args[0][0]['services'], ['Boiler Repair'])
        self.assertNotIn('Water heater repair and replacement', response.get_data(as_text=True))
        # Extra client metadata cannot change campaign/path/template resolution.
        form = submission(['Boiler Repair'])
        form['campaign'] = 'Fake Campaign'
        form['path'] = '/fake'
        form['template'] = 'results.html'
        response = self.client.post(self.path('Boiler Repair'), data=form)
        self.assertEqual(response.status_code, 200)
        self.assertNotIn('Fake Campaign', response.get_data(as_text=True))

    def test_results_controls_all_combinations_and_payloads(self):
        for size in range(1, 7):
            for selected in combinations(SERVICES, size):
                with self.subTest(selected=selected):
                    form = submission(list(selected), state='Massachusetts')
                    with patch('app.generate_landing_pages') as generate:
                        response = self.client.post('/generate', data=form)
                        generate.assert_not_called()
                    self.assertEqual(response.status_code, 200)
                    html = response.get_data(as_text=True)
                    parser = parse(html)
                    self.assertFalse(parser.nested_forms)
                    controls = [f for f in parser.forms if f['attributes']['action'].startswith('/preview/')]
                    self.assertEqual([f['attributes']['action'] for f in controls], [self.path(s) for s in selected])
                    self.assertEqual(len([f for f in parser.forms if f['attributes']['action'] == '/export']), 1)
                    self.assertIn('Download Excel', html)
                    self.assertIn('opens in a new tab', html)
                    self.assertIn('Recommended Landing Page', html)
                    for preview in controls:
                        attrs = preview['attributes']
                        self.assertEqual(attrs['method'], 'post')
                        self.assertEqual(attrs['target'], '_blank')
                        self.assertEqual(attrs['rel'], 'noopener')
                        self.assertNotIn('?', attrs['action'])
                        values = {i['name']: i['value'] for i in preview['inputs'] if i['name'] != 'services'}
                        expected = {key: form[key] for key in form.keys() if key != 'services'}
                        expected.update({
                            'business_phone': '', 'license_information': '',
                            'insurance_information': '', 'additional_business_information': '',
                            'website_url': '',
                        })
                        self.assertEqual(values, expected)
                        self.assertEqual([i['value'] for i in preview['inputs'] if i['name'] == 'services'], list(selected))
                        self.assertTrue(all(i['type'] == 'hidden' for i in preview['inputs']))
                        for service in selected:
                            self.assertIn(f'Preview landing page — {service} (opens in a new tab)', html)

    def test_long_and_malicious_inputs_escape_in_pages_and_hidden_fields(self):
        malicious = '\"><script>alert(1)</script>&\' Long '
        form = submission(SERVICES, business_name=malicious * 40, city=malicious,
                          state=malicious, service_area=malicious * 30)
        for url in ['/generate', self.path('Drain Cleaning')]:
            with self.subTest(url=url):
                response = self.client.post(url, data=form)
                self.assertEqual(response.status_code, 200)
                html = response.get_data(as_text=True)
                parser = parse(html)
                self.assertNotIn('<script>', html)
                self.assertFalse(any(tag == 'script' for tag, attrs in parser.tags))
                self.assertIn('&lt;script&gt;', html)
                self.assertFalse(any('onerror' in attrs or 'onclick' in attrs for tag, attrs in parser.tags))
                if url == '/generate':
                    for rendered_form in parser.forms:
                        values = {i['name']: i['value'] for i in rendered_form['inputs'] if i['name'] != 'services'}
                        for field in ['business_name', 'city', 'state', 'service_area']:
                            self.assertEqual(values[field], form[field].strip())

    def test_independent_business_submissions_do_not_leak(self):
        first = self.client.post(self.path('Emergency Plumbing'), data=submission(business_name='Alpha Service', city='Boston'))
        second = self.client.post(self.path('Emergency Plumbing'), data=submission(
            business_name='Beta Service', city='Chicago', state='Illinois', service_area='Chicago, Evanston'))
        self.assertIn('Alpha Service', first.get_data(as_text=True))
        self.assertNotIn('Beta Service', first.get_data(as_text=True))
        self.assertIn('Beta Service', second.get_data(as_text=True))
        self.assertNotIn('Alpha Service', second.get_data(as_text=True))
        self.assertNotIn('Boston', second.get_data(as_text=True))
        self.assertNotIn('Set-Cookie', first.headers)
        self.assertNotIn('Set-Cookie', second.headers)
        self.assertEqual(self.client.get(self.path('Emergency Plumbing')).headers['Location'], '/')

    def test_export_still_has_eight_sheets_and_no_preview_content(self):
        form = submission(SERVICES)
        self.client.post(self.path('Emergency Plumbing'), data=form)
        response = self.client.post('/export', data=form)
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(response.data))
        self.assertEqual(book.sheetnames, ['Overview', 'Campaign Structure', 'Keywords', 'Negative Keywords',
                                          'Ads', 'Budget', 'Landing Pages', 'Conversion Tracking'])
        data, errors = validate_intake(form)
        plan = generate_campaign(data)
        self.assertEqual(list(book['Landing Pages'].values),
                         [tuple(plan['sections'][4]['headers'])] + [tuple(r) for r in plan['sections'][4]['rows']])
        self.assertIn('Landing-page previews are local only and are not published.',
                      ' '.join(str(r[1]) for r in book['Overview'].values))
        self.assertNotIn('Trust information placeholders', str(list(book['Landing Pages'].values)))


if __name__ == '__main__':
    unittest.main()
