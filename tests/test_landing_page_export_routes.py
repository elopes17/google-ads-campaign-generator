from html.parser import HTMLParser
from io import BytesIO
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import load_workbook

from app import app
from landing_pages import generate_landing_pages
from validation import SERVICES
from test_generator import submission


class FormParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.forms = []
        self.current = None

    def handle_starttag(self, tag, attrs):
        attributes = dict(attrs)
        if tag == 'form':
            self.current = {'attrs': attributes, 'inputs': []}
            self.forms.append(self.current)
        elif tag == 'input' and self.current is not None:
            self.current['inputs'].append(attributes)

    def handle_endtag(self, tag):
        if tag == 'form':
            self.current = None


class LandingPageExportRouteTests(unittest.TestCase):
    def setUp(self):
        self.client = app.test_client()

    def test_valid_download_headers_and_generation(self):
        form = submission(
            ['Emergency Plumbing', 'Boiler Repair'],
            business_name='Fictional Home Services',
        )
        with patch('app.generate_landing_pages', wraps=generate_landing_pages) as generate:
            response = self.client.post('/export/landing-pages', data=form)
            self.assertEqual(generate.call_count, 1)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, 'application/zip')
        self.assertIn('attachment;', response.headers['Content-Disposition'])
        self.assertIn('fictional-home-services-landing-pages.zip', response.headers['Content-Disposition'])
        self.assertNotIn('Set-Cookie', response.headers)
        with ZipFile(BytesIO(response.data)) as archive:
            self.assertEqual(
                archive.namelist(),
                ['index.html', 'emergency-plumber.html', 'boiler-repair.html',
                 'assets/landing_page.css', 'README.txt'],
            )
        self.assertEqual(self.client.get('/export/landing-pages').status_code, 405)

    def test_invalid_intake_retains_values_and_does_not_generate(self):
        cases = [
            {}, submission([]), submission(monthly_budget='NaN'),
            submission(['Unknown']), submission(primary_goal='bad'),
        ]
        for form in cases:
            with self.subTest(form=form):
                with patch('app.generate_landing_pages') as generate:
                    response = self.client.post('/export/landing-pages', data=form)
                    generate.assert_not_called()
                self.assertEqual(response.status_code, 400)
                self.assertEqual(response.mimetype, 'text/html')
                self.assertNotIn('Content-Disposition', response.headers)
                self.assertFalse(response.data.startswith(b'PK'))
                if form:
                    html = response.get_data(as_text=True)
                    self.assertIn(f'value="{form["business_name"]}"', html)
                    self.assertIn(f'value="{form["city"]}"', html)
                    self.assertIn(f'value="{form["monthly_budget"]}"', html)

    def test_results_download_form_payload_and_existing_controls(self):
        form = submission(SERVICES, state='Massachusetts')
        response = self.client.post('/generate', data=form)
        self.assertEqual(response.status_code, 200)
        html = response.get_data(as_text=True)
        self.assertIn('<div class="download-actions" aria-label="Campaign downloads">', html)
        parser = FormParser()
        parser.feed(html)
        downloads = [item for item in parser.forms if item['attrs'].get('action') == '/export/landing-pages']
        self.assertEqual(len(downloads), 1)
        download = downloads[0]
        self.assertEqual(download['attrs'].get('method'), 'post')
        self.assertNotIn('target', download['attrs'])
        values = {item['name']: item['value'] for item in download['inputs'] if item['name'] != 'services'}
        expected = {key: form[key] for key in form.keys() if key != 'services'}
        expected.update({
            'business_phone': '', 'license_information': '',
            'insurance_information': '', 'additional_business_information': '',
            'website_url': '',
        })
        self.assertEqual(values, expected)
        self.assertEqual([item['value'] for item in download['inputs'] if item['name'] == 'services'], SERVICES)
        self.assertTrue(all(item.get('type') == 'hidden' for item in download['inputs']))
        self.assertIn('Download landing pages', html)
        self.assertEqual(len([item for item in parser.forms if item['attrs'].get('action') == '/export']), 1)
        self.assertEqual(len([item for item in parser.forms if item['attrs'].get('action', '').startswith('/preview/')]), 6)

    def test_different_business_downloads_are_isolated(self):
        alpha = self.client.post('/export/landing-pages', data=submission(
            ['Emergency Plumbing'], business_name='Alpha Service', city='Boston', service_area='Boston Area'))
        beta = self.client.post('/export/landing-pages', data=submission(
            ['Emergency Plumbing'], business_name='Beta Service', city='Chicago',
            state='Illinois', service_area='Chicago Area'))
        with ZipFile(BytesIO(alpha.data)) as archive:
            alpha_html = archive.read('emergency-plumber.html').decode('utf-8')
        with ZipFile(BytesIO(beta.data)) as archive:
            beta_html = archive.read('emergency-plumber.html').decode('utf-8')
        self.assertIn('Alpha Service', alpha_html)
        self.assertNotIn('Beta Service', alpha_html)
        self.assertIn('Beta Service', beta_html)
        self.assertNotIn('Alpha Service', beta_html)
        self.assertNotIn('Boston', beta_html)

    def test_browser_preview_and_excel_export_still_work(self):
        form = submission(['Drain Cleaning'], primary_goal='Quote/service requests')
        preview = self.client.post('/preview/drain-cleaning', data=form)
        self.assertEqual(preview.status_code, 200)
        self.assertIn('/static/landing_page.css', preview.get_data(as_text=True))
        self.assertIn('Back to intake form', preview.get_data(as_text=True))
        self.assertNotIn('Back to landing-page index', preview.get_data(as_text=True))

        excel = self.client.post('/export', data=form)
        self.assertEqual(excel.status_code, 200)
        workbook = load_workbook(BytesIO(excel.data))
        self.assertEqual(workbook.sheetnames, [
            'Overview', 'Campaign Structure', 'Keywords', 'Negative Keywords',
            'Ads', 'Budget', 'Landing Pages', 'Conversion Tracking',
        ])
        self.assertEqual({row[1] for row in list(workbook['Landing Pages'].values)[1:]}, {'Drain Cleaning'})


if __name__ == '__main__':
    unittest.main()
