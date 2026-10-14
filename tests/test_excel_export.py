from decimal import Decimal
from io import BytesIO
from itertools import product
from collections import Counter
import unittest
from unittest.mock import patch

from app import app
from generator import generate_campaign
from validation import SERVICES, validate_intake
from test_generator import submission

try:
    from openpyxl import load_workbook
    from excel_export import build_workbook, export_filename, XLSX_MIMETYPE
    HAS_OPENPYXL = True
except ModuleNotFoundError as error:
    if error.name != 'openpyxl':
        raise
    HAS_OPENPYXL = False


@unittest.skipUnless(HAS_OPENPYXL, 'Install requirements.txt to run Excel export tests (openpyxl missing).')
class ExcelExportTests(unittest.TestCase):
    def workbook(self, services=None, **changes):
        data, errors = validate_intake(submission(services, **changes))
        self.assertFalse(errors)
        plan = generate_campaign(data)
        book = load_workbook(build_workbook(data, plan))
        return data, plan, book

    def test_sheet_names_and_formatting(self):
        data, plan, book = self.workbook(SERVICES)
        self.assertEqual(book.sheetnames, ['Overview', 'Campaign Structure', 'Keywords',
            'Negative Keywords', 'Ads', 'Budget', 'Landing Pages', 'Conversion Tracking'])
        for sheet in book:
            self.assertEqual(sheet.freeze_panes, 'A2')
            self.assertTrue(all(c.font.bold for c in sheet[1]))
            self.assertTrue(sheet['A1'].alignment.wrap_text)
            self.assertGreaterEqual(sheet.column_dimensions['A'].width, 16)

    def test_rows_match_generated_plan(self):
        for services in [['Emergency Plumbing'], ['Boiler Repair'], SERVICES]:
            with self.subTest(services=services):
                data, plan, book = self.workbook(services)
                for name, index in [('Campaign Structure', 0), ('Keywords', 1),
                        ('Negative Keywords', 2), ('Landing Pages', 4), ('Conversion Tracking', 5)]:
                    sheet = book[name]
                    self.assertEqual(list(sheet.values),
                        [tuple(plan['sections'][index]['headers'])] + [tuple(r) for r in plan['sections'][index]['rows']])
                expected = []
                for ad in plan['ads']:
                    for kind, key in [('Headline', 'headlines'), ('Description', 'descriptions')]:
                        for number, asset in enumerate(ad[key], 1):
                            expected.append((ad['campaign'], ad['group'], kind, number,
                                             asset['text'], asset['count'], asset['notes']))
                self.assertEqual(list(book['Ads'].values)[1:], expected)

    def test_selected_services_and_asset_counts(self):
        for services in [[s] for s in SERVICES] + [SERVICES]:
            with self.subTest(services=services):
                data, plan, book = self.workbook(services)
                structure = list(book['Campaign Structure'].values)[1:]
                self.assertEqual({r[1] for r in structure}, set(services))
                self.assertEqual({r[0] for r in structure}, {a['campaign'] for a in plan['ads']})
                counts = Counter((r[1], r[2]) for r in list(book['Ads'].values)[1:])
                for service in services:
                    self.assertEqual(counts[service, 'Headline'], 15)
                    self.assertEqual(counts[service, 'Description'], 4)
                for name in ['Keywords', 'Negative Keywords', 'Ads', 'Landing Pages']:
                    self.assertEqual({r[1] for r in list(book[name].values)[1:]}, set(services))

    def test_overview_and_state_formatting(self):
        data, plan, book = self.workbook(SERVICES, state='Massachusetts')
        overview = dict(book['Overview'].values)
        self.assertEqual(overview['Business Name'], data['business_name'])
        self.assertEqual(overview['City'], data['city'])
        self.assertEqual(overview['State'], 'Massachusetts')
        self.assertEqual(overview['Service Area'], data['service_area'])
        self.assertEqual(overview['Selected Services'], ', '.join(SERVICES))
        self.assertEqual(overview['24/7 Availability'], data['emergency_availability'])
        self.assertEqual(overview['Primary Conversion Goal'], data['primary_goal'])
        self.assertEqual([r[1] for r in book['Overview'].values if r[0] == 'Generator Warning'], plan['warnings'])
        keyword_texts = [r[2] for r in list(book['Keywords'].values)[1:]]
        self.assertTrue(any('Boston MA' in text for text in keyword_texts))
        self.assertFalse(any('Massachusetts' in text for text in keyword_texts))
        for row in list(book['Landing Pages'].values)[1:]:
            self.assertIn('Boston, Massachusetts', row[3])

    def test_goals_and_availability(self):
        for available, goal in product(['Yes', 'No'],
                ['Phone calls', 'Quote/service requests', 'Phone calls and quote/service requests']):
            with self.subTest(available=available, goal=goal):
                data, plan, book = self.workbook(SERVICES, emergency_availability=available, primary_goal=goal)
                self.assertEqual(list(book['Conversion Tracking'].values)[1:],
                                 [tuple(r) for r in plan['sections'][5]['rows']])
                text = ' '.join(str(value) for name in ['Ads', 'Keywords', 'Landing Pages']
                                for row in book[name].values for value in row).lower()
                if available == 'No':
                    for claim in ['24/7', '24 hour', 'day or night', 'round-the-clock']:
                        self.assertNotIn(claim, text)
                ads = ' '.join(r[4] for r in list(book['Ads'].values)[1:])
                if goal == 'Quote/service requests':
                    self.assertNotIn('Call', ads)
                elif goal == 'Phone calls':
                    self.assertNotIn('online', ads)
                else:
                    self.assertIn('Call or request service online.', ads)

    def test_numeric_budgets_and_odd_cent_split(self):
        data, plan, book = self.workbook(SERVICES, monthly_budget='2500.01')
        self.assertEqual(book['Overview']['B6'].data_type, 'n')
        self.assertEqual(Decimal(str(book['Overview']['B6'].value)), Decimal('2500.01'))
        for row in [2, 3, 4]:
            for column in ['B', 'C', 'D']:
                self.assertEqual(book['Budget'][f'{column}{row}'].data_type, 'n')
            self.assertEqual(book['Budget'][f'B{row}'].number_format, '$#,##0.00')
        self.assertEqual(Decimal(str(book['Budget']['B2'].value)), Decimal('1250.01'))
        self.assertEqual(Decimal(str(book['Budget']['B3'].value)), Decimal('1250.00'))
        self.assertEqual(Decimal(str(book['Budget']['B4'].value)), Decimal('2500.01'))
        self.assertEqual(book['Budget']['B5'].value, plan['sections'][3]['note'])

    def test_formula_like_inputs_are_literal_text(self):
        for prefix in ['=', '+', '-', '@']:
            with self.subTest(prefix=prefix):
                value = prefix + 'SUM(1,2)'
                data, plan, book = self.workbook(business_name=value, city=value, state=value, service_area=value)
                self.assertEqual(book['Overview']['B2'].value, value)
                self.assertEqual(book['Overview']['B2'].data_type, 's')
                for sheet in book:
                    for row in sheet:
                        for cell in row:
                            self.assertNotEqual(cell.data_type, 'f')
                            if isinstance(cell.value, str):
                                self.assertEqual(cell.data_type, 's')
                self.assertTrue(any(a['text'] == value for a in plan['ads'][0]['headlines']))
                matching = [r for r in list(book['Ads'].values)[1:] if r[4] == value]
                self.assertTrue(matching)

    def test_safe_filenames(self):
        self.assertEqual(export_filename('Fictional Home Services'), 'fictional-home-services-google-ads-plan.xlsx')
        for name in ['../Acme:/\\*?"<>|\r\n', '=SUM(1,2)', 'CON', 'Élite Plumbing', 'A' * 500]:
            filename = export_filename(name)
            self.assertRegex(filename, r'^[a-z0-9-]+\.xlsx$')
            self.assertLessEqual(len(filename), 121)
        for name in ['', '!!!', '水']:
            self.assertEqual(export_filename(name), 'google-ads-campaign-plan.xlsx')

    def test_download_response_and_existing_generator(self):
        with app.test_client() as client:
            form = submission(SERVICES, business_name='Fictional Home Services')
            with patch('app.generate_campaign', wraps=generate_campaign) as generate:
                response = client.post('/export', data=form)
                self.assertEqual(generate.call_count, 1)
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.mimetype, XLSX_MIMETYPE)
            self.assertIn('attachment;', response.headers['Content-Disposition'])
            self.assertIn('fictional-home-services-google-ads-plan.xlsx', response.headers['Content-Disposition'])
            self.assertEqual(len(load_workbook(BytesIO(response.data)).sheetnames), 8)
            self.assertEqual(client.get('/export').status_code, 405)
            html = client.post('/generate', data=form).get_data(as_text=True)
            self.assertIn('action="/export"', html)
            self.assertIn('Download Excel', html)
            for service in SERVICES:
                self.assertIn(f'name="services" value="{service}"', html)


class ExportValidationTests(unittest.TestCase):
    def test_invalid_submissions_do_not_return_workbooks(self):
        with app.test_client() as client:
            for data in [{}, submission([]), submission(monthly_budget='NaN'),
                         submission(['Unknown']), submission(primary_goal='bad')]:
                with self.subTest(data=data):
                    with patch('app.generate_campaign') as generate:
                        response = client.post('/export', data=data)
                        generate.assert_not_called()
                    self.assertEqual(response.status_code, 400)
                    self.assertEqual(response.mimetype, 'text/html')
                    self.assertNotIn('Content-Disposition', response.headers)
                    self.assertFalse(response.data.startswith(b'PK'))
