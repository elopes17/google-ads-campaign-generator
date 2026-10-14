from copy import deepcopy
from io import BytesIO
import unittest

from openpyxl import load_workbook

from business_profile import FIELD_LIMITS, normalize_business_profile
from excel_export import build_workbook
from generator import generate_campaign
from landing_pages import generate_landing_pages
from validation import validate_intake
from test_generator import submission


class BusinessProfileValidationTests(unittest.TestCase):
    def test_missing_and_blank_fields_are_optional(self):
        data, errors = validate_intake(submission())
        self.assertEqual(errors, [])
        for field in FIELD_LIMITS:
            self.assertEqual(data[field], '')
        self.assertEqual(data['business_phone_canonical'], '')

        data, errors = validate_intake(submission(
            business_phone=' \n\t ', license_information=' \n ',
            insurance_information='\t', additional_business_information='  ',
        ))
        self.assertEqual(errors, [])
        for field in FIELD_LIMITS:
            self.assertEqual(data[field], '')

    def test_accepted_phone_formats_have_exact_normalized_values(self):
        for value in [
            '6175550123', '617-555-0123', '(617) 555-0123', '617.555.0123',
            '1-617-555-0123', '+1 617 555 0123',
        ]:
            with self.subTest(value=value):
                profile, errors = normalize_business_profile({'business_phone': value})
                self.assertEqual(errors, [])
                self.assertEqual(profile['business_phone'], '(617) 555-0123')
                self.assertEqual(profile['business_phone_canonical'], '+16175550123')

    def test_invalid_phone_inputs(self):
        invalid = [
            '617555012', '61755501234', '26175550123', '+44 20 7946 0958',
            '617-FIX-PIPES', '617-555-0123 ext 4', 'tel:6175550123',
            'javascript:alert(1)', '617/555/0123', '617+5550123', '++16175550123',
            '617..555..0123', '(617 555-0123', '617) 555-0123',
            '---617-555-0123', '+1...617...555...0123',
            '1 617 555 0123', '617 555 0123',
        ]
        for value in invalid:
            with self.subTest(value=value):
                profile, errors = normalize_business_profile({'business_phone': value})
                self.assertTrue(errors)
                self.assertEqual(profile['business_phone_canonical'], '')

    def test_whitespace_limits_and_unsupported_characters(self):
        form = submission(
            license_information='  License\n\tABC   123 ',
            insurance_information='I' * 160,
            additional_business_information='A' * 240,
        )
        data, errors = validate_intake(form)
        self.assertEqual(errors, [])
        self.assertEqual(data['license_information'], 'License ABC 123')
        self.assertEqual(len(data['insurance_information']), 160)
        self.assertEqual(len(data['additional_business_information']), 240)

        for field, value in [
            ('license_information', 'L' * 161),
            ('insurance_information', 'I' * 161),
            ('additional_business_information', 'A' * 241),
            ('business_phone', '1' * 31),
        ]:
            with self.subTest(field=field):
                data, errors = validate_intake(submission(**{field: value}))
                self.assertTrue(errors)
        for character in ['\x00', '\u200b', '\u202e']:
            with self.subTest(character=repr(character)):
                data, errors = validate_intake(submission(license_information='License' + character))
                self.assertTrue(any('control or formatting' in error for error in errors))

    def test_profile_does_not_change_campaign_output_or_input(self):
        base, errors = validate_intake(submission(['Emergency Plumbing', 'Boiler Repair']))
        self.assertEqual(errors, [])
        profile, errors = validate_intake(submission(
            ['Emergency Plumbing', 'Boiler Repair'], business_phone='+1 617 555 0123',
            license_information='License ABC', insurance_information='Coverage supplied',
            additional_business_information='Family operated',
        ))
        self.assertEqual(errors, [])
        self.assertEqual(generate_campaign(base), generate_campaign(profile))
        before = deepcopy(profile)
        generate_landing_pages(profile)
        self.assertEqual(profile, before)

    def test_page_data_replaces_each_placeholder_independently(self):
        data, errors = validate_intake(submission(
            license_information='License ABC', additional_business_information='Locally operated',
        ))
        self.assertEqual(errors, [])
        page = generate_landing_pages(data)[0]
        items = {item['key']: item for item in page['business_information']}
        self.assertEqual(items['license']['text'], 'License ABC')
        self.assertTrue(items['license']['provided'])
        self.assertTrue(items['insurance']['text'].startswith('Placeholder:'))
        self.assertFalse(items['insurance']['provided'])
        self.assertEqual(items['additional']['text'], 'Locally operated')
        self.assertIn('not independently verified', page['business_information_disclosure'])

    def test_excel_overview_and_formula_protection(self):
        for prefix in ['=', '+', '-', '@']:
            with self.subTest(prefix=prefix):
                value = prefix + 'provided detail'
                data, errors = validate_intake(submission(
                    business_phone='6175550123', license_information=value,
                    insurance_information=value, additional_business_information=value,
                ))
                self.assertEqual(errors, [])
                plan = generate_campaign(data)
                book = load_workbook(build_workbook(data, plan))
                self.assertEqual(len(book.sheetnames), 8)
                overview = {row[0]: row[1] for row in book['Overview'].iter_rows(values_only=True)}
                self.assertEqual(overview['Business Phone'], '(617) 555-0123')
                self.assertEqual(overview['License Information'], value)
                self.assertEqual(overview['Insurance Information'], value)
                self.assertEqual(overview['Additional Business Information'], value)
                for row in range(10, 14):
                    self.assertEqual(book['Overview'].cell(row, 2).data_type, 's')
                self.assertEqual(book['Overview']['B6'].data_type, 'n')
                self.assertEqual(book['Overview']['B6'].number_format, '$#,##0.00')


if __name__ == '__main__':
    unittest.main()
