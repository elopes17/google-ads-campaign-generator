"""Create a local Excel download without changing the generated plan."""
from decimal import Decimal
from io import BytesIO
import re
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

XLSX_MIMETYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
CURRENCY_FORMAT = '$#,##0.00'


def export_filename(business_name):
    ascii_name = unicodedata.normalize('NFKD', business_name).encode('ascii', 'ignore').decode()
    slug = re.sub(r'[^a-z0-9]+', '-', ascii_name.lower()).strip('-')
    slug = slug[:100].rstrip('-')
    if not slug:
        return 'google-ads-campaign-plan.xlsx'
    return f'{slug}-google-ads-plan.xlsx'


def append_row(sheet, values):
    """Force every string to an Excel text cell, including formula-like text."""
    row_number = sheet.max_row + 1 if sheet.max_row > 1 or sheet.cell(1, 1).value is not None else 1
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row_number, column)
        cell.value = value
        if isinstance(value, str):
            # Assignment can initially recognize '=' as a formula. Override it
            # before saving so Excel receives a literal string with no formula.
            cell.data_type = 's'


def format_sheet(sheet):
    sheet.freeze_panes = 'A2'
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for column in range(1, sheet.max_column + 1):
        longest = max(len(str(sheet.cell(row, column).value or '')) for row in range(1, sheet.max_row + 1))
        sheet.column_dimensions[get_column_letter(column)].width = min(65, max(16, longest + 2))


def build_workbook(data, plan):
    workbook = Workbook()
    overview = workbook.active
    overview.title = 'Overview'
    append_row(overview, ['Field', 'Details'])
    for label, value in [
        ('Business Name', data['business_name']), ('City', data['city']),
        ('State', data['state']), ('Service Area', data['service_area']),
        ('Monthly Budget', data['monthly_budget']),
        ('Selected Services', ', '.join(data['services'])),
        ('24/7 Availability', data['emergency_availability']),
        ('Primary Conversion Goal', data['primary_goal']),
        ('Business Phone', data.get('business_phone') or 'Not provided'),
        ('License Information', data.get('license_information') or 'Not provided'),
        ('Insurance Information', data.get('insurance_information') or 'Not provided'),
        ('Additional Business Information', data.get('additional_business_information') or 'Not provided'),
    ]:
        append_row(overview, [label, value])
    overview['B6'].number_format = CURRENCY_FORMAT
    for warning in plan['warnings']:
        append_row(overview, ['Generator Warning', warning])

    sections = {section['id']: section for section in plan['sections']}
    for name, slug in [
        ('Campaign Structure', 'campaign-structure'), ('Keywords', 'keywords'),
        ('Negative Keywords', 'negative-keywords'), ('Ads', None),
        ('Budget', 'budget-allocation'), ('Landing Pages', 'landing-pages'),
        ('Conversion Tracking', 'conversion-tracking'),
    ]:
        sheet = workbook.create_sheet(name)
        if name == 'Ads':
            append_row(sheet, ['Campaign', 'Ad Group', 'Asset Type', 'Asset Number', 'Text', 'Character Count', 'Notes'])
            for ad in plan['ads']:
                for asset_type, key in [('Headline', 'headlines'), ('Description', 'descriptions')]:
                    for number, asset in enumerate(ad[key], start=1):
                        append_row(sheet, [ad['campaign'], ad['group'], asset_type, number,
                                           asset['text'], asset['count'], asset['notes']])
            continue
        section = sections[slug]
        append_row(sheet, section['headers'])
        for original in section['rows']:
            row = list(original)
            if name == 'Budget':
                # Parse the generator's already-rounded display amounts, never
                # recompute allocation here.
                row[1] = Decimal(row[1].replace('$', '').replace(',', ''))
                row[2] = Decimal(row[2].replace('$', '').replace(',', ''))
                row[3] = Decimal(row[3].replace('%', '')) / 100
            append_row(sheet, row)
            if name == 'Budget':
                sheet.cell(sheet.max_row, 2).number_format = CURRENCY_FORMAT
                sheet.cell(sheet.max_row, 3).number_format = CURRENCY_FORMAT
                sheet.cell(sheet.max_row, 4).number_format = '0%'
        if section.get('note'):
            append_row(sheet, ['Note', section['note']])

    for sheet in workbook:
        format_sheet(sheet)
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer
